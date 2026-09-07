#!/usr/bin/env python3
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_test_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0] # Title Layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Apresentação de Teste (Google Slides / PPTX)"
    subtitle.text = "Jewish Educational Resources - Teste do Visualizador"
    
    # Slide 2: Aleph
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    title_shape.text = "Letra Aleph (א)"
    
    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Esta é a primeira letra do Alef-Bet (Alfabeto Hebraico)."
    p = tf.add_paragraph()
    p.text = "Significado simbólico: Força, Líder, Unidade de Deus."
    p.font.size = Pt(20)
    
    # Slide 3: Bet
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    title_shape.text = "Letra Bet (ב)"
    
    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Esta é a segunda letra do Alef-Bet."
    p = tf.add_paragraph()
    p.text = "Significado simbólico: Casa, Criação, Dualidade."
    p.font.size = Pt(20)
    
    output_path = "test-slides.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_test_presentation()
